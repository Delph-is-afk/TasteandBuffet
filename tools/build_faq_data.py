import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


OUTPUT = Path(__file__).resolve().parents[1] / "faq-data.js"
PHONE = "05 58 46 19 48"

ANSWER_OVERRIDES = {
    "Etes-vous ouverts le 31 décembre?": (
        "Oui, un menu et un tarif spécial pour la Saint-Sylvestre seront proposés."
    ),
    "Y a-t-il des fruits de mer ?": "Oui, des fruits de mer sont proposés au buffet.",
    "Y a-t-il des glaces ?": "Oui, des glaces sont proposées au buffet.",
    "Peut-on réserver par WhatsApp ?": (
        "Vous pouvez réserver par WhatsApp au 06 51 86 36 98."
    ),
    "Pouvez-vous réchauffer un biberon ou un petit pot ?": (
        "Oui, le personnel peut réchauffer un biberon ou un petit pot. Un micro-ondes est également à disposition."
    ),
    "Y a-t-il une terrasse ?": "Non, le restaurant n'a pas encore de terrasse extérieure.",
    "Peut-on partager un seul buffet entre plusieurs personnes ?": (
        "Non. Chaque personne doit normalement prendre sa propre formule. Si une personne ne souhaite pas manger, signalez-le à l'avance au personnel."
    ),
    "Les boissons sont-elles servies à table ?": (
        "Les boissons sont servies à table, tandis que la nourriture est en libre-service."
    ),
}


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\u00a0", " ").split())


parser = argparse.ArgumentParser(description="Build faq-data.js from the restaurant FAQ workbook.")
parser.add_argument("source", type=Path, help="Path to the source .xlsx file")
args = parser.parse_args()

workbook = load_workbook(args.source, read_only=True, data_only=True)
sheet = workbook["FAQ recherchée"]
entries = []

for row in sheet.iter_rows(min_row=5, values_only=True):
    category, question, answer, verification, source, verified_at = row
    category = clean(category)
    question = clean(question)
    answer = clean(answer)
    verification = clean(verification)

    if not category or not question:
        continue

    answer = ANSWER_OVERRIDES.get(question, answer)

    needs_call = "confirmer" in verification.lower() or "avis client" in verification.lower()
    if not answer:
        answer = (
            "Cette information n'est pas encore confirmée. "
            f"Le restaurant pourra vous répondre directement au {PHONE}."
        )
        needs_call = True

    entries.append(
        {
            "category": category,
            "question": question,
            "answer": answer,
            "verification": verification,
            "needsCall": needs_call,
        }
    )

payload = json.dumps(entries, ensure_ascii=False, indent=2)
OUTPUT.write_text(
    "// Generated from the restaurant FAQ workbook.\n"
    f"window.TASTE_FAQ = {payload};\n",
    encoding="utf-8",
)

print(f"Generated {len(entries)} FAQ entries in {OUTPUT}")
